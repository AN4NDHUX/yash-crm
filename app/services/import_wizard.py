"""Strict import file readers for guided CRM imports.

Parsing never persists data. The submit endpoint re-parses uploads and enforces
field mappings and authorization before committing any record.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date, datetime
from pathlib import Path

from fastapi import HTTPException

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_IMPORT_ROWS = 100_000
MAX_IMPORT_COLUMNS = 200
ENCODINGS = {
    "auto": "utf-8-sig", "utf-8": "utf-8-sig", "utf-16": "utf-16",
    "iso-8859-1": "iso-8859-1", "iso-8859-2": "iso-8859-2",
    "iso-8859-8": "iso-8859-8", "iso-8859-9": "iso-8859-9",
    "iso-8859-11": "iso-8859-11", "gb2312": "gb2312", "gbk": "gbk",
    "big5": "big5", "shift_jis": "shift_jis",
}


def _value(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _decode(raw: bytes, charset: str) -> str:
    codec = ENCODINGS.get(charset.lower())
    if not codec:
        raise HTTPException(422, "Unsupported character encoding")
    try:
        return raw.decode(codec)
    except UnicodeError as exc:
        raise HTTPException(422, "File cannot be decoded with selected charset") from exc


def _cards(content: str):
    rows = []
    fields = {}
    # RFC 6350 folding: a continuation line begins with a space/tab.
    lines = re.sub(r"\r?\n[ \t]", "", content).splitlines()
    for line in lines:
        line = line.strip()
        if line.upper() == "BEGIN:VCARD":
            fields = {}
        elif line.upper() == "END:VCARD":
            name = fields.get("FN") or fields.get("N", "").replace(";", " ").strip()
            parts = name.rsplit(" ", 1)
            rows.append({
                "First Name": parts[0] if len(parts) > 1 else name,
                "Last Name": parts[-1] if len(parts) > 1 else name,
                "Full Name": name,
                "Phone": fields.get("TEL", ""),
                "Email": fields.get("EMAIL", ""),
                "Company": fields.get("ORG", ""),
                "Job Title": fields.get("TITLE", ""),
                "Notes": fields.get("NOTE", ""),
            })
            fields = {}
        elif ":" in line:
            key, value = line.split(":", 1)
            key = key.split(";", 1)[0].upper()
            if key in {"FN", "N", "TEL", "EMAIL", "ORG", "TITLE", "NOTE"} and key not in fields:
                fields[key] = value.replace("\\n", " ").replace("\\,", ",").strip()
    return rows


def parse_import_file(filename: str, raw: bytes, charset: str = "auto") -> dict:
    filename = Path(filename or "").name
    ext = Path(filename).suffix.lower()
    if ext not in {".csv", ".xlsx", ".xls", ".vcf"}:
        raise HTTPException(422, "Supported file types: CSV, XLSX, XLS and VCF")
    if not raw or len(raw) > MAX_FILE_BYTES:
        raise HTTPException(413 if len(raw) > MAX_FILE_BYTES else 422, "File must be nonempty and no larger than 25 MB")
    if ext == ".vcf":
        records = _cards(_decode(raw, charset))
        headers = ["First Name", "Last Name", "Full Name", "Phone", "Email", "Company", "Job Title", "Notes"]
    elif ext == ".csv":
        try:
            reader = csv.reader(io.StringIO(_decode(raw, charset), newline=""))
            headers = next(reader, [])
            records = []
            for number, values in enumerate(reader, start=1):
                if number > MAX_IMPORT_ROWS:
                    raise HTTPException(413, "Only 100,000 records can be imported")
                records.append(values)
        except HTTPException:
            raise
        except (csv.Error, ValueError) as exc:
            raise HTTPException(422, "Invalid CSV file") from exc
    elif ext == ".xlsx":
        try:
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                    raise HTTPException(413, "Workbook exceeds safe uncompressed size")
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            try:
                rows = wb.active.iter_rows(values_only=True)
                headers = next(rows, [])
                records = []
                for index, row in enumerate(rows, start=1):
                    if index > MAX_IMPORT_ROWS:
                        raise HTTPException(413, "Only 100,000 records can be imported")
                    records.append(row)
            finally:
                wb.close()
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, "Invalid XLSX workbook") from exc
    else:
        try:
            import xlrd
            book = xlrd.open_workbook(file_contents=raw)
            sheet = book.sheet_by_index(0)
            if sheet.nrows > MAX_IMPORT_ROWS + 1:
                raise HTTPException(413, "Only 100,000 records can be imported")
            headers = sheet.row_values(0) if sheet.nrows else []
            records = [sheet.row_values(index) for index in range(1, sheet.nrows)]
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, "Invalid XLS workbook") from exc
    if ext == ".vcf":
        headers = list(headers)
    else:
        headers = [_value(h) for h in headers]
    if not headers or len(headers) > MAX_IMPORT_COLUMNS or any(not name for name in headers):
        raise HTTPException(422, "File needs 1–200 named header columns")
    if len(set(h.casefold() for h in headers)) != len(headers):
        raise HTTPException(422, "Duplicate column headers are not supported")
    if len(records) > MAX_IMPORT_ROWS:
        raise HTTPException(413, "Only 100,000 records can be imported")
    if ext == ".vcf":
        result = records
    else:
        result = [{head: _value(values[i]) if i < len(values) else "" for i, head in enumerate(headers)}
                  for values in records if any(_value(v) for v in values)]
    if not result:
        raise HTTPException(422, "The file contains no records")
    return {"filename": filename, "columns": headers, "rows": result, "count": len(result), "sample": result[:3]}
